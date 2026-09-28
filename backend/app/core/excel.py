"""Geração da planilha Excel formatada, com fórmulas vivas.

Mesmo com o painel na web, o cliente final quer o arquivo — e as taxas vão como
FÓRMULA do Excel (não número congelado), então ele pode editar a meta ou uma
linha e ver o recálculo.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .metricas import TAXAS_COMERCIAIS, TAXAS_CRUZADAS, TAXAS_META
from .pipeline import Resultado

AZUL, AZUL_CLARO, CINZA = "1F3864", "D9E2F3", "F2F2F2"
F_TIT = Font(name="Arial", bold=True, size=14, color=AZUL)
F_SUB = Font(name="Arial", italic=True, size=9, color="595959")
F_CAB = Font(name="Arial", bold=True, color="FFFFFF", size=10)
F_TXT = Font(name="Arial", size=10)
F_TOT = Font(name="Arial", bold=True, size=10)
FILL_CAB = PatternFill("solid", fgColor=AZUL)
FILL_TOT = PatternFill("solid", fgColor=AZUL_CLARO)
FILL_BLOCO = PatternFill("solid", fgColor=CINZA)
FILL_OK = PatternFill("solid", fgColor="E2EFDA")
FILL_RUIM = PatternFill("solid", fgColor="FCE4D6")
BORDA = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
FMT = {
    "brl": '"R$" #,##0.00;-"R$" #,##0.00;"-"',
    "pct": '0.0%;-0.0%;"-"',
    "int": '#,##0;-#,##0;"-"',
    "x": '0.00"x";-0.00"x";"-"',
    "dec": '#,##0.0;-#,##0.0;"-"',
    "data": "dd/mm/yyyy",
}


def formato_coluna(nome: str, serie: pd.Series | None = None) -> str | None:
    n = nome.lower()
    if "(r$)" in n:
        return FMT["brl"]
    if "(%)" in n:
        return FMT["pct"]
    if "(x)" in n:
        return FMT["x"]
    if "(dias)" in n or "dias" in n:
        return FMT["dec"]
    if serie is not None and pd.api.types.is_datetime64_any_dtype(serie):
        return FMT["data"]
    if serie is not None and pd.api.types.is_numeric_dtype(serie):
        return FMT["int"] if (serie.dropna() % 1 == 0).all() else FMT["dec"]
    return None


def valor_excel(v):
    if v is None or (isinstance(v, float) and np.isnan(v)) or v is pd.NA or v is pd.NaT:
        return None
    if isinstance(v, np.generic):
        return v.item()
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    return v


def _cabecalho(ws, titulo: str, nome_mes: str, subtitulo: str | None = None) -> int:
    ws["A1"] = titulo
    ws["A1"].font = F_TIT
    ws["A2"] = subtitulo or f"Período: {nome_mes}  •  Gerado em {datetime.now():%d/%m/%Y %H:%M}"
    ws["A2"].font = F_SUB
    ws.sheet_view.showGridLines = False
    return 4


def _escrever_tabela(ws, df, linha, titulo=None, taxas=None, total=True, nao_somar=(), col=1) -> int:
    """Escreve df a partir de 'linha'. Colunas de taxa viram FÓRMULAS. Retorna próxima linha livre."""
    if titulo:
        ws.cell(linha, col, titulo).font = Font(name="Arial", bold=True, size=11, color=AZUL)
        linha += 1
    if df is None or df.empty:
        ws.cell(linha, col, "Sem dados para este bloco.").font = F_SUB
        return linha + 2
    taxas = [
        t
        for t in (taxas or [])
        if all(c in df for c in t["num"]) and (t["den"] == "__TOTAL__" or all(c in df for c in t["den"]))
    ]
    df = df.drop(columns=[t["nome"] for t in taxas if t["nome"] in df])
    cols = list(df.columns) + [t["nome"] for t in taxas]
    L = {c: get_column_letter(col + i) for i, c in enumerate(cols)}
    for i, c in enumerate(cols):
        cel = ws.cell(linha, col + i, c)
        cel.font, cel.fill, cel.border = F_CAB, FILL_CAB, BORDA
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[linha].height = 32
    cab, ini, fim = linha, linha + 1, linha + len(df)
    lin_tot = fim + 1

    def formula(t, rr):
        num = "+".join(f"{L[c]}{rr}" for c in t["num"])
        den = f"{L[t['num'][0]]}${lin_tot}" if t["den"] == "__TOTAL__" else "+".join(f"{L[c]}{rr}" for c in t["den"])
        mult = f"*{t['mult']}" if t.get("mult") else ""
        return f"=IFERROR(({num})/({den}){mult},0)"

    for j, (_, row) in enumerate(df.iterrows()):
        rr = ini + j
        for i, c in enumerate(df.columns):
            cel = ws.cell(rr, col + i, valor_excel(row[c]))
            cel.font, cel.border = F_TXT, BORDA
            f = formato_coluna(c, df[c])
            if f:
                cel.number_format = f
            if j % 2:
                cel.fill = FILL_BLOCO
        for t in taxas:
            cel = ws.cell(rr, col + cols.index(t["nome"]), formula(t, rr))
            cel.font, cel.border = F_TXT, BORDA
            cel.number_format = formato_coluna(t["nome"]) or FMT["pct"]
            if j % 2:
                cel.fill = FILL_BLOCO
    if total:
        for i, c in enumerate(cols):
            cel = ws.cell(lin_tot, col + i)
            cel.font, cel.fill, cel.border = F_TOT, FILL_TOT, BORDA
            if i == 0:
                cel.value = "TOTAL"
            elif c in df and pd.api.types.is_numeric_dtype(df[c]) and c not in nao_somar:
                cel.value = f"=SUM({L[c]}{ini}:{L[c]}{fim})"
                cel.number_format = formato_coluna(c, df[c]) or FMT["int"]
            elif c not in df:
                t = next(t for t in taxas if t["nome"] == c)
                cel.value = formula(t, lin_tot)
                cel.number_format = formato_coluna(c) or FMT["pct"]
    ultima = lin_tot if total else fim
    if not ws.auto_filter.ref:
        ws.auto_filter.ref = f"{get_column_letter(col)}{cab}:{get_column_letter(col + len(cols) - 1)}{fim}"
    return ultima + 2


def _ajustar_larguras(ws, minimo=10, maximo=55) -> None:
    larg: dict[str, int] = {}
    for row in ws.iter_rows(min_row=3):
        for c in row:
            if c.value is not None and not str(c.value).startswith("="):
                larg[c.column_letter] = max(larg.get(c.column_letter, 0), len(str(c.value)))
    for letra, w in larg.items():
        ws.column_dimensions[letra].width = max(minimo, min(maximo, w * 0.95 + 2))


def _envolver_texto(ws, min_row=5) -> None:
    for row in ws.iter_rows(min_row=min_row):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


# ====================================================================== #
def gerar_excel(res: Resultado, destino: str | Path | None = None) -> bytes:
    """Monta a planilha completa. Devolve os bytes e, se `destino` vier, também salva."""
    t = res.tabelas
    nome_mes = res.nome_mes
    wb = Workbook()

    # ---------------- 1. Resumo executivo ----------------
    ws = wb.active
    ws.title = "Resumo Executivo"
    lin = _cabecalho(ws, f"Assertividade Comercial — {nome_mes}", nome_mes)
    for i, h in enumerate(["Indicador", "Valor", "Como é calculado", "Meta", "Situação"], 1):
        c = ws.cell(lin, i, h)
        c.font, c.fill, c.border = F_CAB, FILL_CAB, BORDA
    lin += 1
    REF: dict[str, str] = {}
    cursor = lin
    for _, chaves in res.blocos:  # reserva a linha de cada KPI antes de escrever as fórmulas
        cursor += 1
        for k in chaves:
            REF[k] = f"B{cursor}"
            cursor += 1
        cursor += 1
    import re as _re

    for titulo, chaves in res.blocos:
        c = ws.cell(lin, 1, titulo)
        c.font = Font(name="Arial", bold=True, color=AZUL, size=11)
        for i in range(1, 6):
            ws.cell(lin, i).fill = FILL_TOT
        lin += 1
        for k in chaves:
            dados = res.kpis[k]
            ws.cell(lin, 1, dados["rotulo"]).font = Font(name="Arial", size=10, bold=(k == "assert"))
            if dados["formula"]:
                v = "=IFERROR(" + _re.sub(r"\{(\w+)\}", lambda mm: REF[mm.group(1)], dados["formula"][1:]) + ",0)"
            else:
                v = dados["valor"]
            cel = ws.cell(lin, 2, valor_excel(v))
            cel.number_format = FMT[dados["formato"]]
            cel.font = Font(name="Arial", size=10, bold=True)
            ws.cell(lin, 3, dados["explicacao"]).font = F_SUB
            if dados["meta_tipo"]:
                cm = ws.cell(lin, 4, dados["meta_valor"])
                cm.number_format = FMT[dados["formato"]]
                cm.font = Font(name="Arial", size=10, color="0000FF")
                cond = f"B{lin}>=D{lin}" if dados["meta_tipo"] == "min" else f"AND(B{lin}<=D{lin},B{lin}>0)"
                ws.cell(lin, 5, f'=IF({cond},"✅ Dentro da meta","⚠️ Fora da meta")').font = F_TXT
                for i in (1, 2, 5):
                    ws.cell(lin, i).fill = FILL_OK if dados["dentro_da_meta"] else FILL_RUIM
            for i in range(1, 6):
                ws.cell(lin, i).border = BORDA
            lin += 1
        lin += 1
    ws.cell(lin, 1, "Metas em azul são editáveis — a coluna Situação recalcula sozinha.").font = F_SUB
    for letra, w in zip("ABCDE", [44, 18, 62, 14, 20]):
        ws.column_dimensions[letra].width = w
    ws.freeze_panes = "A5"

    # ---------------- 2. Diagnóstico ----------------
    ws = wb.create_sheet("Diagnóstico")
    lin = _cabecalho(ws, "Diagnóstico automático — o que os números estão dizendo", nome_mes)
    _escrever_tabela(ws, res.diagnostico, lin, total=False)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 130
    _envolver_texto(ws)

    # ---------------- 3. Funil ----------------
    funil = t.get("funil", pd.DataFrame())
    if not funil.empty:
        ws = wb.create_sheet("Funil Completo")
        lin = _cabecalho(ws, "Funil completo: do anúncio ao cliente (somente tráfego pago)", nome_mes)
        ws.cell(lin, 1, "Investimento (R$)").font = F_TOT
        ws.cell(lin, 2, res.kpi("inv")).number_format = FMT["brl"]
        lin += 2
        f_ini = lin + 1
        cabs = [
            "Etapa", "Quantidade", "Conversão da etapa anterior (%)",
            "Conversão desde impressões (%)", "Custo por etapa (R$)",
        ]
        for i, h in enumerate(cabs, 1):
            c = ws.cell(lin, i, h)
            c.font, c.fill, c.border = F_CAB, FILL_CAB, BORDA
            c.alignment = Alignment(wrap_text=True, horizontal="center")
        for j, row in funil.reset_index(drop=True).iterrows():
            rr = f_ini + j
            ws.cell(rr, 1, row["Etapa"])
            ws.cell(rr, 2, valor_excel(row["Quantidade"])).number_format = FMT["int"]
            ws.cell(rr, 3, "" if j == 0 else f"=IFERROR(B{rr}/B{rr - 1},0)").number_format = FMT["pct"]
            ws.cell(rr, 4, f"=IFERROR(B{rr}/$B${f_ini},0)").number_format = "0.000%"
            ws.cell(rr, 5, f"=IFERROR($B$4/B{rr},0)").number_format = FMT["brl"]
            for i in range(1, 6):
                ws.cell(rr, i).border, ws.cell(rr, i).font = BORDA, F_TXT
        graf = BarChart()
        graf.type = "bar"
        graf.title = "Funil (etapas após o clique)"
        graf.style = 10
        graf.add_data(Reference(ws, min_col=2, min_row=f_ini + 1, max_row=f_ini + 5), titles_from_data=False)
        graf.set_categories(Reference(ws, min_col=1, min_row=f_ini + 1, max_row=f_ini + 5))
        graf.legend = None
        graf.y_axis.majorGridlines = None
        graf.x_axis.scaling.orientation = "maxMin"
        graf.height, graf.width = 8, 18
        ws.add_chart(graf, f"A{f_ini + 8}")
        ws.cell(
            f_ini + 7, 1,
            "Obs.: impressões ficam fora do gráfico porque a escala esconderia as outras etapas.",
        ).font = F_SUB
        for letra, w in zip("ABCDE", [40, 16, 20, 20, 18]):
            ws.column_dimensions[letra].width = w

    # ---------------- 4. Meta ----------------
    ws = wb.create_sheet("Meta - Campanhas")
    lin = _cabecalho(ws, "Meta Ads — desempenho por campanha", nome_mes)
    lin = _escrever_tabela(ws, t.get("meta_campanhas"), lin, taxas=TAXAS_META, nao_somar=["Tipo de resultado"])
    ws.cell(
        lin, 1,
        "Alcance e frequência no TOTAL são somas entre campanhas — podem conter pessoas repetidas.",
    ).font = F_SUB
    _ajustar_larguras(ws)
    ws.freeze_panes = "B5"

    for nome, chave, titulo, congelar in [
        ("Meta - Conjuntos", "meta_conjuntos", "Meta Ads — desempenho por conjunto de anúncios", "C5"),
        ("Meta - Anúncios", "meta_anuncios", "Meta Ads — desempenho por anúncio (criativo)", "D5"),
    ]:
        if chave in t and not t[chave].empty:
            ws = wb.create_sheet(nome)
            lin = _cabecalho(ws, titulo, nome_mes)
            _escrever_tabela(ws, t[chave], lin, taxas=TAXAS_META)
            _ajustar_larguras(ws)
            ws.freeze_panes = congelar

    # ---------------- 5. Campanha x Vendas ----------------
    ws = wb.create_sheet("Campanha x Vendas")
    lin = _cabecalho(ws, "Qual campanha realmente vende — Meta cruzada com as reuniões", nome_mes)
    if "campanha_x_vendas" not in t or t["campanha_x_vendas"].empty:
        ws.cell(
            lin, 1,
            "A planilha de reuniões não tem a coluna 'Campanha'. Registre de qual campanha veio cada lead "
            "para liberar esta aba.",
        ).font = F_SUB
    else:
        _escrever_tabela(ws, t["campanha_x_vendas"], lin, taxas=TAXAS_CRUZADAS)
    _ajustar_larguras(ws)
    ws.freeze_panes = "B5"

    # ---------------- 6. Comercial ----------------
    ws = wb.create_sheet("Status das Reuniões")
    lin = _cabecalho(ws, "Distribuição dos status das reuniões", nome_mes)
    por_status = t.get("por_status", pd.DataFrame())
    lin = _escrever_tabela(
        ws,
        por_status.drop(columns="% do total (%)", errors="ignore"),
        lin,
        "Status padronizado",
        taxas=[{"nome": "% do total (%)", "num": ["Quantidade"], "den": "__TOTAL__"}],
    )
    lin = _escrever_tabela(ws, t.get("mapa_status"), lin, "Como cada texto da planilha foi classificado")
    _escrever_tabela(ws, t.get("status_por_responsavel"), lin, "Status × responsável")
    _ajustar_larguras(ws)

    for nome, chave, titulo in [
        ("Por Responsável", "por_responsavel", "Assertividade por responsável (closer / vendedor)"),
        ("Por Origem", "por_origem", "Assertividade por origem do lead"),
        ("Por Produto", "por_produto", "Assertividade por produto / serviço"),
        ("Por Marcação", "por_marcacao", "Assertividade por marcação no nome do cliente"),
        ("Por Semana", "por_semana", "Evolução semana a semana"),
        ("Por Dia da Semana", "por_dia", "Assertividade por dia da semana da reunião"),
    ]:
        if chave not in t or t[chave].empty:
            continue
        ws = wb.create_sheet(nome)
        lin = _cabecalho(ws, titulo, nome_mes)
        _escrever_tabela(ws, t[chave], lin, taxas=TAXAS_COMERCIAIS)
        _ajustar_larguras(ws, maximo=24)
        ws.freeze_panes = "B5"

    ws = wb.create_sheet("Ciclo de Venda")
    lin = _cabecalho(ws, "Ciclo de venda — quanto tempo leva para fechar", nome_mes)
    lin = _escrever_tabela(ws, t.get("tempo_agendamento"), lin, "Tempo entre o agendamento e a reunião", total=False)
    lin = _escrever_tabela(
        ws, t.get("ciclo"), lin, "Reunião → fechamento (precisa da coluna 'Data de fechamento')", total=False
    )
    lin = _escrever_tabela(ws, t.get("ciclo_por_responsavel"), lin, "Reunião → fechamento por responsável",
                           total=False)
    if "fechamentos" in t:
        _escrever_tabela(ws, t["fechamentos"], lin, "Todos os fechamentos do mês", nao_somar=["Dias até fechar"])
    _ajustar_larguras(ws)

    ws = wb.create_sheet("Motivos de Perda")
    lin = _cabecalho(ws, "Por que os clientes não fecharam", nome_mes)
    lin = _escrever_tabela(
        ws,
        t.get("motivos_perda", pd.DataFrame()).drop(columns="% das perdas (%)", errors="ignore"),
        lin,
        taxas=[{"nome": "% das perdas (%)", "num": ["Quantidade"], "den": "__TOTAL__"}],
    )
    if "lista_perdas" in t:
        _escrever_tabela(ws, t["lista_perdas"], lin, "Lista de perdas")
    _ajustar_larguras(ws)

    ws = wb.create_sheet("Pipeline em Aberto")
    lin = _cabecalho(ws, "Oportunidades em aberto — ordenadas por temperatura", nome_mes)
    _escrever_tabela(ws, t.get("pipeline"), lin, nao_somar=["Dias em aberto"], total=res.tem_valor)
    _ajustar_larguras(ws, maximo=70)
    ws.freeze_panes = "B5"
    _envolver_texto(ws)

    ws = wb.create_sheet("Observações e Objeções")
    lin = _cabecalho(ws, "O que as observações revelam (objeções, sinais de compra e de perda)", nome_mes)
    lin = _escrever_tabela(ws, t.get("sinais_resumo"), lin, "Sinais mais frequentes", total=False)
    lin = _escrever_tabela(ws, t.get("sinais_por_status"), lin, "Sinal × status da reunião")
    _escrever_tabela(ws, t.get("observacoes"), lin, "Todas as observações, cliente a cliente", total=False)
    _ajustar_larguras(ws, maximo=80)
    _envolver_texto(ws)

    ws = wb.create_sheet("No-show e Remarcações")
    lin = _cabecalho(ws, "Reuniões que não aconteceram — lista para recuperar", nome_mes)
    _escrever_tabela(ws, t.get("noshow"), lin, total=False)
    _ajustar_larguras(ws)

    # ---------------- 7. Bases tratadas ----------------
    ws = wb.create_sheet("Base Reuniões (tratada)")
    lin = _cabecalho(ws, "Base de reuniões do mês, já padronizada (use para conferir qualquer número)", nome_mes)
    _escrever_tabela(ws, t.get("base_reunioes"), lin, total=False)
    _ajustar_larguras(ws, maximo=35)
    ws.freeze_panes = "A5"

    ws = wb.create_sheet("Base Meta (tratada)")
    lin = _cabecalho(ws, "Base da Meta do período, já padronizada", nome_mes)
    _escrever_tabela(ws, t.get("base_meta"), lin, total=False)
    _ajustar_larguras(ws, maximo=40)

    # ---------------- 8. Qualidade + glossário ----------------
    ws = wb.create_sheet("Qualidade dos Dados")
    lin = _cabecalho(ws, "Qualidade dos dados — o que pode estar distorcendo os números", nome_mes)
    lin = _escrever_tabela(ws, res.resumo_qualidade, lin, "Resumo", total=False)
    lin = _escrever_tabela(ws, res.qualidade, lin, "Detalhe linha a linha", total=False)
    _escrever_tabela(ws, res.mapeamento, lin, "Como as colunas dos seus arquivos foram reconhecidas", total=False)
    _ajustar_larguras(ws)

    ws = wb.create_sheet("Glossário")
    lin = _cabecalho(ws, "Glossário — definição de cada indicador", nome_mes)
    extras = [
        ("Reunião realizada",
         "Status Fechado, Perdido, Em negociação/Follow-up ou Reunião feita (a conversa aconteceu)"),
        ("Temperatura (pipeline)",
         "🔥 Quente = Follow-up ou promessa de fechamento na observação; ❄️ Fria = sinal de perda; 🌤️ Morna = o resto"),
        ("Marcação", "Texto depois do hífen no nome do cliente (ex.: 'Wagner - Parceiro' → Parceiro)"),
        ("Status 'Agendado'", "Reunião futura; não entra em nenhuma taxa até ser atualizada"),
        ("Tráfego pago", "Origem contém: " + ", ".join(res.config.origens_trafego_pago)),
        ("Status padronizado", "Texto livre da planilha convertido pelas regras do sistema"),
    ]
    gl = pd.DataFrame(
        [(d["rotulo"], d["explicacao"]) for d in res.kpis.values()] + extras,
        columns=["Indicador", "Definição"],
    )
    _escrever_tabela(ws, gl, lin, total=False)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 110

    wb.calculation.fullCalcOnLoad = True
    buffer = BytesIO()
    wb.save(buffer)
    dados = buffer.getvalue()
    if destino:
        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_bytes(dados)
    return dados
