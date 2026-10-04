"""Planilha-modelo para quem usa HubSpot.

Quem já tem CRM não quer preencher planilha: quer exportar. Então este arquivo
não é um formulário em branco com outro cabeçalho, é um **gabarito de
exportação**. Ele diz exatamente quais propriedades marcar na tela de export do
HubSpot, na ordem, e traz três linhas de exemplo com o conteúdo que cada coluna
deve ter — para a pessoa conferir se o que saiu do CRM dela se parece com isso
antes de subir.

Por que um modelo separado em vez de só documentar os apelidos: o usuário de
HubSpot não sabe que o Neriah chama "Deal Name" de cliente. Ele sabe o nome que
vê na tela dele. O modelo fala a língua dele e some com a tradução.

Os cabeçalhos aqui são os nomes de propriedade do HubSpot, e são os mesmos que
`config_analise.COLUNAS_REUNIOES` reconhece. Mudar um sem mudar o outro quebra
o reconhecimento em silêncio, e é por isso que `tests/test_modelo_hubspot.py`
gera esta planilha, preenche e manda para o motor: se o par sair de sincronia,
o teste cai.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .modelo import (
    BORDA,
    F_CAB,
    F_EX,
    F_SEC,
    F_SUB,
    F_TIT,
    F_TXT,
    FILL_CAB,
    FILL_OBR,
    FILL_ZEBRA,
    _aba_metas,
)

ABA = "Negócios (HubSpot)"

# (propriedade no HubSpot, obrigatória?, o que o Neriah faz com ela, exemplo)
PROPRIEDADES = [
    ("Deal Name", True, "Identifica o cliente. É a linha da análise.", "Padaria do Zé"),
    ("Deal Stage", True,
     "Vira a etapa do funil. Closed Won conta como fechado, Closed Lost como perdido.",
     "Closed Won"),
    ("Close Date", True,
     "Situa o negócio no mês. Sem data o Neriah não sabe qual mês analisar.",
     "2026-09-05"),
    ("Amount", False,
     "Destrava receita, ticket médio, ROAS e o valor do pipeline em aberto.", "2500.00"),
    ("Deal owner", False,
     "Destrava a comparação entre vendedores: quem fecha mais e quem perde onde.",
     "Marina Alves"),
    ("Original Traffic Source", False,
     "Separa o que veio de anúncio do que veio de indicação. É o que torna o CAC real.",
     "Paid Social"),
    ("Create Date", False,
     "Destrava o tempo de ciclo: quantos dias da entrada até o fechamento.",
     "2026-09-01"),
    ("Notes", False,
     "De onde saem as objeções e a temperatura do pipeline. A coluna mais subestimada.",
     "Achou caro, comparou com a contabilidade atual"),
    ("Closed Lost Reason", False,
     "Agrupa os motivos de perda. Sem ela, 'por que perdemos' fica sem resposta.",
     "Preço"),
]

# As etapas padrão de um pipeline de vendas do HubSpot, na ordem do funil.
ETAPAS_HUBSPOT = [
    "Appointment Scheduled",
    "Qualified To Buy",
    "Presentation Scheduled",
    "Decision Maker Bought-In",
    "Contract Sent",
    "Closed Won",
    "Closed Lost",
]

EXEMPLOS = [
    ["Padaria do Zé", "Closed Won", "2026-09-05", "2500.00", "Marina Alves", "Paid Social",
     "2026-09-01", "Fechou na call, indicou um vizinho", ""],
    ["Mercado Sul", "Closed Lost", "2026-09-06", "1800.00", "Marina Alves", "Paid Social",
     "2026-09-02", "Comparou com a contabilidade atual e achou caro", "Preço"],
    ["Bar do João", "Presentation Scheduled", "2026-09-11", "3200.00", "Rafael Pinto",
     "Organic Search", "2026-09-04", "Pediu o contrato para analisar, retorna dia 12", ""],
]

PASSOS = [
    "No HubSpot, abra CRM → Negócios (Deals).",
    "Clique em Ações → Exportar visualização, ou Exportar no topo da lista.",
    "Em 'Propriedades a exportar', marque as nove da aba ao lado. As três primeiras "
    "são obrigatórias; cada uma das outras acende um indicador.",
    "Escolha o formato CSV ou XLSX e confirme. O HubSpot manda o arquivo por e-mail.",
    "Suba esse arquivo no Neriah junto com o relatório de campanhas da Meta.",
]


def _aba_negocios(wb: Workbook) -> None:
    ws = wb.create_sheet(ABA, 0)
    ws["A1"] = "Negócios exportados do HubSpot"
    ws["A1"].font = F_TIT
    ws["A2"] = (
        "Este é o formato que o Neriah espera. Você não precisa preencher nada aqui: exporte do "
        "HubSpot com as propriedades da aba 'Como exportar' e o arquivo já sai assim. As três "
        "linhas abaixo são exemplo, para você comparar com o que saiu do seu CRM."
    )
    ws["A2"].font = F_SUB
    ws.sheet_view.showGridLines = False

    linha_cab = 4
    for i, (prop, obrigatoria, _destrava, _ex) in enumerate(PROPRIEDADES, start=1):
        cel = ws.cell(linha_cab, i, prop)
        cel.font, cel.fill, cel.border = F_CAB, FILL_CAB, BORDA
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(16, min(34, len(prop) + 10))
    ws.row_dimensions[linha_cab].height = 30

    for j, exemplo in enumerate(EXEMPLOS):
        for i, valor in enumerate(exemplo, start=1):
            cel = ws.cell(linha_cab + 1 + j, i, valor)
            cel.font, cel.border = F_EX, BORDA
            if j % 2:
                cel.fill = FILL_ZEBRA

    primeira_livre = linha_cab + 1 + len(EXEMPLOS)
    ultima = primeira_livre + 300
    for r in range(primeira_livre, ultima):
        for i in range(1, len(PROPRIEDADES) + 1):
            ws.cell(r, i).border = BORDA
            ws.cell(r, i).font = F_TXT

    # Lista suspensa na etapa, com as etapas padrão do HubSpot. Serve para quem
    # preenche à mão em vez de exportar.
    col_etapa = next(i for i, (p, *_) in enumerate(PROPRIEDADES, start=1) if p == "Deal Stage")
    letra = get_column_letter(col_etapa)
    dv = DataValidation(
        type="list", formula1='"' + ",".join(ETAPAS_HUBSPOT) + '"',
        allow_blank=True, showDropDown=False,
    )
    dv.error = "Escolha uma das etapas padrão do HubSpot."
    dv.errorTitle = "Etapa inválida"
    ws.add_data_validation(dv)
    dv.add(f"{letra}{linha_cab + 1}:{letra}{ultima}")

    for i, (_prop, obrigatoria, *_r) in enumerate(PROPRIEDADES, start=1):
        if obrigatoria:
            for r in range(primeira_livre, primeira_livre + 40):
                ws.cell(r, i).fill = FILL_OBR

    ws.freeze_panes = f"A{linha_cab + 1}"
    ws.auto_filter.ref = f"A{linha_cab}:{get_column_letter(len(PROPRIEDADES))}{ultima}"
    ws.cell(ultima + 2, 1, "As três primeiras linhas são exemplo — apague antes de usar.").font = F_SUB


def _aba_como_exportar(wb: Workbook) -> None:
    ws = wb.create_sheet("Como exportar")
    ws["A1"] = "Como tirar isto do HubSpot"
    ws["A1"].font = F_TIT
    ws["A2"] = "Cinco passos. Leva menos de dois minutos."
    ws["A2"].font = F_SUB
    ws.sheet_view.showGridLines = False

    lin = 4
    for i, passo in enumerate(PASSOS, start=1):
        cel = ws.cell(lin, 1, f"{i}. {passo}")
        cel.font = F_TXT
        cel.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=lin, start_column=1, end_row=lin, end_column=3)
        ws.row_dimensions[lin].height = 30
        lin += 1

    lin += 1
    ws.cell(lin, 1, "As propriedades a marcar").font = F_SEC
    lin += 1
    for i, titulo in enumerate(["Propriedade no HubSpot", "Precisa?", "O que destrava"], 1):
        c = ws.cell(lin, i, titulo)
        c.font, c.fill, c.border = F_CAB, FILL_CAB, BORDA
        c.alignment = Alignment(horizontal="center", vertical="center")
    lin += 1
    for j, (prop, obrigatoria, destrava, _ex) in enumerate(PROPRIEDADES):
        valores = [prop, "OBRIGATÓRIA" if obrigatoria else "opcional", destrava]
        for i, v in enumerate(valores, 1):
            cel = ws.cell(lin, i, v)
            cel.font, cel.border = F_TXT, BORDA
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            if obrigatoria:
                cel.fill = FILL_OBR
            elif j % 2:
                cel.fill = FILL_ZEBRA
        ws.row_dimensions[lin].height = 30
        lin += 1

    for letra, largura in zip("ABC", [30, 16, 74]):
        ws.column_dimensions[letra].width = largura

    lin += 2
    ws.cell(lin, 1, "Se o nome da sua propriedade for outro").font = F_SEC
    lin += 1
    cel = ws.cell(
        lin, 1,
        "Pode ser. Quem personalizou o CRM às vezes renomeia as propriedades, e o Neriah "
        "reconhece vários nomes para o mesmo campo: 'Deal Name', 'Deal - Title', 'Cliente' e "
        "'Negócio' chegam todos no mesmo lugar. O que não dá é faltar a propriedade. Se o Neriah "
        "não reconhecer alguma, a aba 'Qualidade dos dados' da análise mostra exatamente qual "
        "coluna ele não entendeu.",
    )
    cel.font, cel.alignment = F_TXT, Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=lin, start_column=1, end_row=lin + 3, end_column=3)

    lin += 5
    ws.cell(lin, 1, "O que o HubSpot não exporta, e por que importa").font = F_SEC
    lin += 1
    for aviso in [
        "Não existe etapa de no-show no pipeline padrão: quem faltou fica em "
        "'Appointment Scheduled', igual a quem ainda vai ser atendido. O comparecimento sai "
        "100% por falta de registro, e o Neriah avisa isso na análise em vez de fingir.",
        "A exportação de negócios não traz a campanha que originou cada um. Sem ela, o Neriah "
        "pontua as campanhas por custo de lead e clique; com ela, passa a pontuar por cliente "
        "fechado, que é o que decide onde cortar verba. Se você usa 'Original Traffic Source "
        "Drill-Down 1' para guardar a campanha, marque essa propriedade também.",
    ]:
        cel = ws.cell(lin, 1, aviso)
        cel.font = F_TXT
        cel.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=lin, start_column=1, end_row=lin + 1, end_column=3)
        ws.row_dimensions[lin].height = 30
        lin += 3


def gerar_planilha_hubspot(destino: str | Path | None = None) -> bytes:
    """Monta o gabarito de exportação do HubSpot. Mesma aba de Metas do modelo padrão."""
    wb = Workbook()
    wb.remove(wb.active)
    _aba_negocios(wb)
    _aba_metas(wb)
    _aba_como_exportar(wb)

    dados = BytesIO()
    wb.save(dados)
    conteudo = dados.getvalue()
    if destino:
        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_bytes(conteudo)
    return conteudo
